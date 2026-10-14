import os
import re
import logging
import zipfile
import pdfplumber
from docx import Document

MAX_PAGES_READ = 2


def extract_text(file):
    ext = os.path.splitext(file.filename)[1].lower()
    try:
        if ext == ".pdf":
            text, page_count = _from_pdf(file.stream)
        else:
            text, page_count = _from_docx(file.stream)
        return text, page_count, None
    except Exception:
        logging.exception("Could not read uploaded file")
        return None, None, "Cannot read this file. It may be corrupted or password protected."


def _from_pdf(stream):
    parts = []
    with pdfplumber.open(stream) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages[:MAX_PAGES_READ]:
            parts.append(page.extract_text() or "")
    return "\n".join(parts), page_count


def _from_docx(stream):
    doc = Document(stream)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    stream.seek(0)
    return "\n".join(lines), _docx_pages(stream)


def _docx_pages(stream):
    try:
        with zipfile.ZipFile(stream) as z:
            xml = z.read("docProps/app.xml").decode("utf-8", "ignore")
        match = re.search(r"<Pages>(\d+)</Pages>", xml)
        return int(match.group(1)) if match else None
    except Exception:
        return None